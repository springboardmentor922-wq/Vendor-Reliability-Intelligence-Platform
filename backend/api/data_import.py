"""Spreadsheet upload: preview, import, template download and history.

Flow on screen: the user drops a workbook, ``/preview`` validates every row
inside a transaction that is rolled back and returns what *would* happen,
then ``/commit`` runs the same import for real. After a commit the platform
recomputes reliability scores and runs the alert sweep, so every dashboard,
ranking and notification reflects the new data immediately.
"""

from __future__ import annotations

import io
import json
import os
import time
import uuid
from datetime import date, timedelta
from typing import Optional

import pandas as pd
from fastapi import APIRouter, Depends, File, HTTPException, UploadFile, status
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from pydantic import BaseModel, Field
from sqlalchemy import text
from sqlalchemy.orm import Session

from config import settings
from database import get_db
from deps import require_roles
from models import (
    ContractStatus,
    InvoiceStatus,
    NotificationType,
    PurchaseOrderStatus,
    User,
    UserRole,
    VendorCategory,
    VendorStatus,
)
from services import data_import as engine
from services.events import log_activity, notify_roles

router = APIRouter(prefix="/data-import", tags=["Data Import"])

require_importer = require_roles(
    UserRole.ADMINISTRATOR,
    UserRole.PROCUREMENT_MANAGER,
    UserRole.SUPPLY_CHAIN_MANAGER,
)

IMPORT_DIR = os.path.join(settings.UPLOAD_DIR, "imports")
os.makedirs(IMPORT_DIR, exist_ok=True)

MAX_IMPORT_BYTES = 150 * 1024 * 1024
STAGED_TTL_SECONDS = 6 * 3600


class CommitRequest(BaseModel):
    token: str = Field(min_length=8, max_length=64)
    dataco_orders: Optional[int] = Field(default=None, ge=100, le=20000)
    dataco_months: Optional[int] = Field(default=None, ge=6, le=36)


# --------------------------------------------------------------------------
# Staging helpers
# --------------------------------------------------------------------------

def _staged_path(token: str) -> tuple[str, dict]:
    meta_path = os.path.join(IMPORT_DIR, f"{token}.json")
    if not token.isalnum() or not os.path.exists(meta_path):
        raise HTTPException(status_code=404, detail="That upload has expired - please upload the file again")
    with open(meta_path) as handle:
        meta = json.load(handle)
    return os.path.join(IMPORT_DIR, meta["stored_name"]), meta


def _cleanup_old() -> None:
    cutoff = time.time() - STAGED_TTL_SECONDS
    for name in os.listdir(IMPORT_DIR):
        path = os.path.join(IMPORT_DIR, name)
        try:
            if os.path.getmtime(path) < cutoff:
                os.remove(path)
        except OSError:
            pass


def _dataco_summary(frame: pd.DataFrame) -> dict:
    columns = {engine.norm(c): c for c in frame.columns}
    order_col = columns.get("order_id")
    date_col = columns.get("order_date_dateorders")
    summary = {
        "rows": int(len(frame)),
        "orders": int(frame[order_col].nunique()) if order_col else None,
        "first_order": None,
        "last_order": None,
    }
    if date_col:
        dates = pd.to_datetime(frame[date_col], errors="coerce").dropna()
        if not dates.empty:
            summary["first_order"] = dates.min().date().isoformat()
            summary["last_order"] = dates.max().date().isoformat()
    return summary


def _record_import(db: Session, user: User, report: engine.ImportReport, status_text: str) -> None:
    totals = report.totals
    db.execute(
        text(
            """
            INSERT INTO data_imports
              (file_name, file_format, mode, status, rows_read, rows_created,
               rows_updated, rows_skipped, summary, imported_by)
            VALUES (:f, :fmt, :mode, :status, :rows, :created, :updated, :skipped, :summary, :user)
            """
        ),
        {
            "f": report.file_name[:255],
            "fmt": report.file_format,
            "mode": report.mode,
            "status": status_text,
            "rows": totals["rows"],
            "created": totals["created"],
            "updated": totals["updated"],
            "skipped": totals["skipped"],
            "summary": json.dumps({
                "sheets": [{"sheet": s.sheet, "entity": s.entity, "created": s.created, "updated": s.updated, "skipped": s.skipped} for s in report.sheets],
                "errors": report.error_count,
                "post_processing": report.post_processing,
            }),
            "user": user.id,
        },
    )


def _post_process(db: Session) -> dict:
    """Rescore every vendor and raise alerts for the newly imported data."""

    from services.notifier import run_alert_sweep
    from services.reliability import recalculate_all

    started = time.perf_counter()
    scored = recalculate_all(db)
    db.commit()

    # The sweep is idempotent (one alert per event per user), so the number
    # of genuinely new alerts is the change in the notification count.
    before = db.execute(text("SELECT COUNT(*) FROM notifications")).scalar() or 0
    alerts = run_alert_sweep(db)
    db.commit()
    after = db.execute(text("SELECT COUNT(*) FROM notifications")).scalar() or 0

    return {
        "vendors_rescored": len(scored),
        "alerts_raised": int(after - before),
        "alert_conditions_checked": alerts,
        "seconds": round(time.perf_counter() - started, 2),
    }


# --------------------------------------------------------------------------
# Routes
# --------------------------------------------------------------------------

@router.post("/preview")
async def preview(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    current_user: User = Depends(require_importer),
):
    _cleanup_old()

    content = await file.read()
    name = os.path.basename(file.filename or "upload.xlsx")

    if not content:
        raise HTTPException(status_code=400, detail="The uploaded file is empty")
    if len(content) > MAX_IMPORT_BYTES:
        raise HTTPException(status_code=413, detail="Imports are limited to 150 MB")

    try:
        file_format, frames = engine.read_frames(content, name)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc))
    except Exception as exc:  # noqa: BLE001 - corrupt workbook, wrong format
        raise HTTPException(status_code=400, detail=f"The file could not be read: {exc}")

    token = uuid.uuid4().hex
    stored_name = f"{token}{os.path.splitext(name)[1].lower()}"
    with open(os.path.join(IMPORT_DIR, stored_name), "wb") as handle:
        handle.write(content)
    with open(os.path.join(IMPORT_DIR, f"{token}.json"), "w") as handle:
        json.dump({"file_name": name, "stored_name": stored_name, "user_id": current_user.id}, handle)

    report = engine.ImportReport(file_name=name, file_format=file_format)

    dataco_sheet = engine.is_dataco(frames)
    if dataco_sheet:
        report.mode = "dataco"
        summary = _dataco_summary(frames[dataco_sheet])
        report.notes.append(
            "Recognised the DataCo supply-chain dataset. Importing it rebuilds the dataset-derived "
            "order history (purchase orders, deliveries, invoices, quality evaluations) through the "
            "same ETL the command-line loader uses, then rescores every vendor."
        )
        payload = report.as_dict()
        payload["token"] = token
        payload["dataco"] = summary
        return payload

    try:
        engine.run_workbook(db, current_user, frames, report)
    finally:
        db.rollback()

    payload = report.as_dict()
    payload["token"] = token
    return payload


@router.post("/commit")
def commit(
    body: CommitRequest,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_importer),
):
    path, meta = _staged_path(body.token)

    with open(path, "rb") as handle:
        content = handle.read()

    file_format, frames = engine.read_frames(content, meta["file_name"])
    report = engine.ImportReport(file_name=meta["file_name"], file_format=file_format)

    dataco_sheet = engine.is_dataco(frames)

    if dataco_sheet:
        report.mode = "dataco"
        frame = frames[dataco_sheet]
        csv_path = os.path.join(IMPORT_DIR, f"{body.token}.dataco.csv")
        frame.to_csv(csv_path, index=False, encoding="latin-1", errors="replace")

        from etl.load_dataco import SOURCE_TAG
        from etl.load_dataco import load as load_dataco
        from services.reliability import backfill_history

        summary = _dataco_summary(frame)
        orders = body.dataco_orders or min(6000, summary["orders"] or 6000)
        months = body.dataco_months or 24

        started = time.perf_counter()
        load_dataco(orders=orders, months=months, dataset_path=csv_path)
        db.expire_all()

        counts = db.execute(text(
            "SELECT (SELECT COUNT(*) FROM purchase_orders WHERE source_ref LIKE :tag),"
            "       (SELECT COUNT(*) FROM vendor_performance),"
            "       (SELECT COUNT(*) FROM invoices)"
        ), {"tag": f"{SOURCE_TAG}%"}).one()

        backfill_history(db, months=12)
        db.commit()
        post = _post_process(db)
        post["history_backfilled_months"] = 12
        post["etl_seconds"] = round(time.perf_counter() - started, 1)

        sheet = engine.SheetResult(sheet=dataco_sheet, entity="purchase_orders", rows=summary["rows"], created=int(counts[0]))
        report.sheets.append(sheet)
        report.post_processing = post
        report.committed = True
        report.notes.append(
            f"Loaded {int(counts[0]):,} dataset purchase orders spanning {months} months; "
            f"{int(counts[1]):,} performance evaluations and {int(counts[2]):,} invoices now on record."
        )
        try:
            os.remove(csv_path)
        except OSError:
            pass
    else:
        engine.run_workbook(db, current_user, frames, report)
        db.commit()
        report.committed = True
        report.post_processing = _post_process(db)

    totals = report.totals

    log_activity(
        db, current_user.id, "DataImport", None, "Spreadsheet Imported",
        f"{report.file_name}: {totals['created']} created, {totals['updated']} updated, "
        f"{totals['skipped']} skipped across {len(report.sheets)} sheet(s)",
    )

    notify_roles(
        db,
        [UserRole.ADMINISTRATOR, UserRole.PROCUREMENT_MANAGER],
        NotificationType.SYSTEM,
        "Spreadsheet data imported",
        f"{current_user.name} imported {report.file_name}: {totals['created']} new and "
        f"{totals['updated']} updated record(s). Dashboards and reliability scores were refreshed.",
        link="/data-import",
        exclude_user_id=current_user.id,
    )

    _record_import(db, current_user, report, "Completed with errors" if report.error_count else "Completed")
    db.commit()

    for name in (os.path.basename(path), f"{body.token}.json"):
        try:
            os.remove(os.path.join(IMPORT_DIR, name))
        except OSError:
            pass

    return report.as_dict()


@router.get("/history")
def history(
    db: Session = Depends(get_db),
    current_user: User = Depends(require_importer),
):
    rows = db.execute(
        text(
            """
            SELECT d.id, d.file_name, d.file_format, d.mode, d.status, d.rows_read,
                   d.rows_created, d.rows_updated, d.rows_skipped, d.summary,
                   d.created_at, u.name AS imported_by
            FROM data_imports d
            LEFT JOIN users u ON u.id = d.imported_by
            ORDER BY d.id DESC
            LIMIT 25
            """
        )
    ).mappings().all()

    return [
        {
            **{k: v for k, v in row.items() if k not in ("summary", "created_at")},
            "summary": json.loads(row["summary"]) if row["summary"] else None,
            "created_at": row["created_at"].isoformat() if row["created_at"] else None,
        }
        for row in rows
    ]


# --------------------------------------------------------------------------
# Template
# --------------------------------------------------------------------------

TEMPLATE_SHEETS: list[tuple[str, list[str], list[list]]] = [
    (
        "Vendors",
        ["vendor_code", "vendor_name", "category", "status", "contact_person", "email", "phone", "website",
         "address", "city", "state", "postal_code", "country", "tax_id", "registration_number",
         "company_type", "year_established", "products_services"],
        [
            ["", "Apex Components Pvt Ltd", "Raw Material Suppliers", "Approved", "Ravi Kumar", "ravi@apexcomponents.in",
             "+91 98480 12345", "apexcomponents.in", "Plot 12, IDA Jeedimetla", "Hyderabad", "Telangana", "500055",
             "India", "36AABCA1234F1Z5", "U27100TG2012PTC080011", "Private Limited", 2012, "Steel coils, aluminium sheets"],
            ["", "BlueLine Freight Services", "Logistics Partners", "Approved", "Anita Rao", "ops@bluelinefreight.in",
             "+91 90000 45678", "bluelinefreight.in", "Survey 88, Shamshabad", "Hyderabad", "Telangana", "501218",
             "India", "36AAFCB5678K1Z2", "U60200TG2015PTC099120", "Private Limited", 2015, "FTL/LTL road freight, warehousing"],
        ],
    ),
    (
        "Certifications",
        ["vendor", "certification_name", "issuing_authority", "certificate_number", "issue_date", "expiry_date"],
        [
            ["Apex Components Pvt Ltd", "ISO 9001:2015", "Bureau Veritas", "BV-QMS-44821", "2024-04-01", "2027-03-31"],
            ["BlueLine Freight Services", "ISO 14001:2015", "TUV SUD", "TUV-EMS-10233", "2023-10-15", "2026-10-14"],
        ],
    ),
    (
        "Procurement Requests",
        ["request_number", "item", "category", "quantity", "unit", "estimated_cost", "currency", "required_date",
         "priority", "department", "justification", "status", "vendor"],
        [
            ["", "Cold-rolled steel coils", "Raw Material Suppliers", 40, "Tonnes", 52000, "USD", "2026-11-15",
             "High", "Production", "Q4 production plan", "Approved", "Apex Components Pvt Ltd"],
        ],
    ),
    (
        "Purchase Orders",
        ["po_number", "vendor", "title", "department", "order_date", "expected_delivery", "actual_delivery",
         "status", "currency", "payment_terms", "item_name", "quantity", "unit", "unit_price", "tax_rate", "quality_rating"],
        [
            ["IMP-PO-1001", "Apex Components Pvt Ltd", "Steel coils - Aug batch", "Production", "2026-08-02", "2026-08-20",
             "2026-08-19", "Completed", "USD", "Net 30", "Cold-rolled steel coil 1.2mm", 20, "Tonnes", 1250, 18, 4.5],
            ["IMP-PO-1001", "Apex Components Pvt Ltd", "Steel coils - Aug batch", "Production", "2026-08-02", "2026-08-20",
             "2026-08-19", "Completed", "USD", "Net 30", "Aluminium sheet 2mm", 5, "Tonnes", 2100, 18, 4.5],
            ["IMP-PO-1002", "BlueLine Freight Services", "Inbound freight - Sep", "Logistics", "2026-09-05", "2026-09-12",
             "2026-09-15", "Delivered", "USD", "Net 15", "FTL Hyderabad-Chennai", 3, "Trips", 900, 18, 3.8],
            ["IMP-PO-1003", "BlueLine Freight Services", "Inbound freight - Oct", "Logistics", "2026-09-22", "2026-10-06",
             "", "Ordered", "USD", "Net 15", "LTL consolidated shipment", 2, "Trips", 650, 18, ""],
        ],
    ),
    (
        "Invoices",
        ["invoice_number", "po_number", "invoice_date", "due_date", "amount", "tax_amount", "status", "payment_date"],
        [
            ["IMP-INV-5001", "IMP-PO-1001", "2026-08-21", "2026-09-20", 35500, 6390, "Paid", "2026-09-18"],
            ["IMP-INV-5002", "IMP-PO-1002", "2026-09-16", "2026-10-01", 2700, 486, "Pending", ""],
        ],
    ),
    (
        "Contracts",
        ["contract_number", "vendor", "title", "contract_type", "start_date", "expiry_date", "contract_value",
         "currency", "status", "compliance_status", "auto_renew", "renewal_notice_days"],
        [
            ["", "Apex Components Pvt Ltd", "Annual steel supply agreement", "Supply Agreement", "2026-01-01",
             "2026-12-31", 420000, "USD", "", "Compliant", "No", 45],
            ["", "BlueLine Freight Services", "Freight rate contract", "Service Agreement", "2025-11-01",
             "2026-10-31", 96000, "USD", "", "Compliant", "Yes", 30],
        ],
    ),
    (
        "Performance",
        ["vendor", "po_number", "evaluation_date", "quality_rating", "service_rating", "response_time", "issue_resolution_time", "remarks"],
        [
            ["BlueLine Freight Services", "IMP-PO-1002", "2026-09-16", 3.8, 4.0, 6, 30, "Late by 3 days, goods intact"],
        ],
    ),
]

INSTRUCTIONS = [
    ("How this workbook imports", ""),
    ("1", "Keep one record type per sheet. Sheet and column names are matched loosely, so your own export usually works."),
    ("2", "Vendors are imported first, so other sheets can refer to a vendor by its name or vendor_code."),
    ("3", "Leave vendor_code, po_number, contract_number and invoice_number blank to have them generated."),
    ("4", "Re-importing a row with an existing code or number updates that record instead of duplicating it."),
    ("5", "Purchase Orders: repeat the same po_number on several rows to add several line items to one order."),
    ("6", "Purchase Orders: an actual_delivery date marks the order delivered; on-time vs delayed is computed from expected_delivery."),
    ("7", "tax_rate is a percentage (18 = 18%). quality_rating is 0-5 and feeds the Product Quality reliability factor."),
    ("8", "Imported vendors default to status Approved (existing, onboarded suppliers). New suppliers should use the registration form."),
    ("9", "Purchase orders can only be raised against Approved vendors, matching the application's business rules."),
    ("10", "After import every vendor is rescored and the alert sweep runs, so dashboards update immediately."),
    ("", ""),
    ("Categories", ", ".join(VendorCategory.ALL)),
    ("Vendor status", ", ".join(VendorStatus.ALL)),
    ("PO status", ", ".join(PurchaseOrderStatus.ALL)),
    ("Contract status", ", ".join(ContractStatus.ALL) + " (blank = derived from dates)"),
    ("Invoice status", ", ".join(InvoiceStatus.ALL)),
    ("", ""),
    ("DataCo dataset", "You can also upload DataCoSupplyChainDataset.csv (or an .xlsx copy) - it is recognised automatically and loaded through the dataset ETL."),
]


@router.get("/template")
def template(current_user: User = Depends(require_importer)):
    workbook = Workbook()
    info = workbook.active
    info.title = "Instructions"
    info.column_dimensions["A"].width = 22
    info.column_dimensions["B"].width = 120
    for row in INSTRUCTIONS:
        info.append(list(row))
    info["A1"].font = Font(bold=True, size=14)

    header_fill = PatternFill("solid", fgColor="4F46E5")

    for title, headers, rows in TEMPLATE_SHEETS:
        sheet = workbook.create_sheet(title)
        sheet.append(headers)
        for row in rows:
            sheet.append(row)
        for index, header in enumerate(headers, start=1):
            cell = sheet.cell(row=1, column=index)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = header_fill
            cell.alignment = Alignment(vertical="center")
            sheet.column_dimensions[get_column_letter(index)].width = max(14, len(header) + 4)
        sheet.freeze_panes = "A2"

        def validate(column: str, values: list[str]) -> None:
            if column not in headers:
                return
            letter = get_column_letter(headers.index(column) + 1)
            rule = DataValidation(type="list", formula1=f'"{",".join(values)}"', allow_blank=True)
            sheet.add_data_validation(rule)
            rule.add(f"{letter}2:{letter}2000")

        if title in ("Vendors", "Procurement Requests"):
            validate("category", VendorCategory.ALL)
        if title == "Vendors":
            validate("status", VendorStatus.ALL)
        if title == "Purchase Orders":
            validate("status", PurchaseOrderStatus.ALL)
        if title == "Invoices":
            validate("status", InvoiceStatus.ALL)
        if title == "Contracts":
            validate("status", ContractStatus.ALL)

    buffer = io.BytesIO()
    workbook.save(buffer)
    buffer.seek(0)

    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="VendorIQ_Import_Template.xlsx"'},
    )
