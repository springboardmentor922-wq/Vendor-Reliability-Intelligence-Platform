"""
VendorIQ Excel Report Generator
Generates formatted Excel reports using openpyxl.
"""
import io
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from backend.database import get_db
from backend.calculations import compute_vendor_performance_and_reliability

def generate_vendor_performance_excel() -> bytes:
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute("SELECT id, vendor_code, company_name, category, status FROM vendors ORDER BY id ASC")
    vendors = cursor.fetchall()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Vendor Intelligence Matrix"

    # Header styling
    header_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    border_thin = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    headers = [
        "Vendor Code", "Company Name", "Category", "Status",
        "Reliability Score", "Risk Level", "Procurement Recommendation",
        "Delivery Rate (%)", "Quality Score", "Communication Score",
        "Compliance Score", "Purchase History Score", "Issue Resolution Score",
        "Total Orders", "Completed Orders", "Avg Response Time (h)"
    ]

    ws.append(headers)
    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = align_center

    for v in vendors:
        m = compute_vendor_performance_and_reliability(v["id"], conn)
        factors = m["factors"]
        row = [
            v["vendor_code"],
            v["company_name"],
            v["category"],
            v["status"],
            m["overall_reliability"],
            m["risk_level"],
            m["recommendation"],
            m["delivery_rate"],
            factors["product_quality"]["score"],
            factors["communication_efficiency"]["score"],
            factors["contract_compliance"]["score"],
            factors["purchase_history"]["score"],
            factors["issue_resolution"]["score"],
            m["total_orders"],
            m["completed_orders"],
            m["avg_response_hours"]
        ]
        ws.append(row)

    conn.close()

    # Apply borders and auto-fit column widths
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=len(headers)):
        for cell in row:
            cell.border = border_thin
            if isinstance(cell.value, (int, float)):
                cell.alignment = align_center
            else:
                cell.alignment = align_left

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()

def generate_procurement_excel() -> bytes:
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute("""
        SELECT po.po_number, v.vendor_code, v.company_name, v.category,
               po.order_date, po.expected_delivery_date, po.actual_delivery_date,
               po.total_amount, po.status, po.shipping_mode,
               inv.invoice_number, inv.payment_status
        FROM purchase_orders po
        JOIN vendors v ON po.vendor_id = v.id
        LEFT JOIN invoices inv ON inv.po_id = po.id
        ORDER BY po.order_date DESC
    """)
    rows = cursor.fetchall()
    conn.close()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Procurement & Purchase Orders"

    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

    headers = [
        "PO Number", "Vendor Code", "Vendor Name", "Category",
        "Order Date", "Expected Delivery", "Actual Delivery",
        "Total Amount ($)", "Order Status", "Shipping Mode",
        "Invoice Number", "Payment Status"
    ]
    ws.append(headers)
    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    for r in rows:
        ws.append([
            r["po_number"],
            r["vendor_code"],
            r["company_name"],
            r["category"],
            str(r["order_date"]),
            str(r["expected_delivery_date"]),
            str(r["actual_delivery_date"]) if r["actual_delivery_date"] else "Pending",
            r["total_amount"],
            r["status"],
            r["shipping_mode"],
            r["invoice_number"] or "—",
            r["payment_status"] or "—"
        ])

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 14)

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()
