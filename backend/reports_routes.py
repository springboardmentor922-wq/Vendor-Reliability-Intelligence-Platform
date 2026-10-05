from datetime import datetime
from io import BytesIO

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import (
    SimpleDocTemplate,
    Table,
    TableStyle,
    Paragraph,
)
from sqlalchemy.orm import Session

from database import get_db
from auth import get_current_user
from models import (
    User,
    Vendor,
    VendorPerformance,
    ProcurementRequest,
    PurchaseOrder,
    Supplier,
    Contract,
    Invoice,
    Payment,
)


router = APIRouter(
    prefix="/reports",
    tags=["Reports"]
)


# ============================================================
# ROLE -> REPORT ACCESS
# ============================================================

ROLE_REPORTS = {
    "administrator": [
        "procurement",
        "purchase-orders",
        "vendor-performance",
        "compliance",
        "contracts",
    ],
    "procurement_manager": [
        "procurement",
        "purchase-orders",
        "vendor-performance",
        "contracts",
    ],
    "supply_chain_manager": [
        "purchase-orders",
        "vendor-performance",
    ],
    "finance_officer": [
        "invoices",
        "payments",
    ],
    "auditor": [
        "procurement",
        "purchase-orders",
        "vendor-performance",
        "compliance",
        "contracts",
        "invoices",
        "payments",
    ],
    "vendor": [
        "procurement",
        "purchase-orders",
        "vendor-performance",
        "contracts",
        "invoices",
        "payments",
    ],
}


ALL_REPORT_TYPES = [
    "procurement",
    "purchase-orders",
    "vendor-performance",
    "compliance",
    "contracts",
    "invoices",
    "payments",
]


# ============================================================
# COMMON HELPERS
# ============================================================

def format_date(value):
    if value is None:
        return "—"

    if isinstance(value, datetime):
        return value.strftime("%d %b %Y")

    try:
        return value.strftime("%d %b %Y")
    except Exception:
        return str(value)


def format_datetime(value):
    if value is None:
        return "—"

    if isinstance(value, datetime):
        return value.strftime("%d %b %Y %H:%M")

    try:
        return value.strftime("%d %b %Y %H:%M")
    except Exception:
        return str(value)


def format_amount(value):
    if value is None:
        return "₹0.00"

    try:
        return f"₹{float(value):,.2f}"
    except Exception:
        return str(value)


def format_status(value):
    if value is None:
        return "—"

    text = str(value)

    return (
        text
        .replace("_", " ")
        .replace("-", " ")
        .title()
    )


def format_score(value):
    if value is None:
        return "No data"

    try:
        return f"{float(value):.2f}"
    except Exception:
        return str(value)


def safe_value(value):
    if value is None:
        return "—"

    return str(value)


# ============================================================
# PERFORMANCE CALCULATION
# ============================================================

PERFORMANCE_WEIGHTS = {
    "delivery": 40,
    "quality": 25,
    "communication": 15,
    "compliance": 20,
}


def calculate_weighted_performance(
    performance,
    reliability=None
):
    """
    Calculate Overall Performance using available detailed metrics.

    Weights:
        Delivery      = 40%
        Quality       = 25%
        Communication = 15%
        Compliance    = 20%

    Missing metrics are excluded from the calculation.
    They are NOT treated as zero.

    If no detailed metrics are available, the vendor's
    reliability score is used as Overall Performance.
    """

    if performance:
        metrics = []

        delivery = getattr(
            performance,
            "delivery_score",
            None
        )

        quality = getattr(
            performance,
            "quality_score",
            None
        )

        communication = getattr(
            performance,
            "communication_score",
            None
        )

        compliance = getattr(
            performance,
            "compliance_score",
            None
        )

        # Delivery
        if delivery is not None:
            try:
                metrics.append(
                    (
                        float(delivery),
                        PERFORMANCE_WEIGHTS["delivery"]
                    )
                )
            except (TypeError, ValueError):
                pass

        # Quality
        if quality is not None:
            try:
                metrics.append(
                    (
                        float(quality),
                        PERFORMANCE_WEIGHTS["quality"]
                    )
                )
            except (TypeError, ValueError):
                pass

        # Communication
        if communication is not None:
            try:
                metrics.append(
                    (
                        float(communication),
                        PERFORMANCE_WEIGHTS["communication"]
                    )
                )
            except (TypeError, ValueError):
                pass

        # Compliance
        if compliance is not None:
            try:
                metrics.append(
                    (
                        float(compliance),
                        PERFORMANCE_WEIGHTS["compliance"]
                    )
                )
            except (TypeError, ValueError):
                pass

        # Calculate weighted score using only available metrics
        if metrics:
            total_weight = sum(
                weight
                for _, weight in metrics
            )

            weighted_total = sum(
                score * weight
                for score, weight in metrics
            )

            if total_weight > 0:
                return round(
                    weighted_total / total_weight,
                    2
                )

    # ========================================================
    # FALLBACK TO VENDOR RELIABILITY
    # ========================================================

    if reliability is not None:
        try:
            return round(
                float(reliability),
                2
            )
        except (TypeError, ValueError):
            pass

    return None


def get_latest_vendor_performance(
    db,
    vendor_id
):
    return (
        db.query(VendorPerformance)
        .filter(
            VendorPerformance.vendor_id == vendor_id
        )
        .order_by(
            VendorPerformance.id.desc()
        )
        .first()
    )


# ============================================================
# ACCESS CONTROL
# ============================================================

def get_allowed_reports(current_user):
    role = str(current_user.role)

    return ROLE_REPORTS.get(
        role,
        []
    )


def check_report_access(
    current_user,
    report_type=None
):
    role = str(current_user.role)

    if role not in ROLE_REPORTS:
        raise HTTPException(
            status_code=403,
            detail="You are not authorized to access reports"
        )

    if report_type is not None:
        if report_type not in ROLE_REPORTS[role]:
            raise HTTPException(
                status_code=403,
                detail=(
                    f"Your role does not have access "
                    f"to the {report_type} report"
                )
            )

    return True


# ============================================================
# VENDOR SECURITY
# ============================================================

def get_current_vendor(
    db: Session,
    current_user
):
    vendor = (
        db.query(Vendor)
        .filter(
            Vendor.user_id == current_user.id
        )
        .first()
    )

    if not vendor:
        raise HTTPException(
            status_code=404,
            detail="Vendor profile not found"
        )

    return vendor


# ============================================================
# INTERNAL VENDOR PERFORMANCE REPORT
# ============================================================

def get_vendor_performance_report(db):
    vendors = (
        db.query(Vendor)
        .order_by(
            Vendor.id.asc()
        )
        .all()
    )

    report = []

    for vendor in vendors:
        performance = get_latest_vendor_performance(
            db,
            vendor.id
        )

        delivery = (
            getattr(
                performance,
                "delivery_score",
                None
            )
            if performance
            else None
        )

        quality = (
            getattr(
                performance,
                "quality_score",
                None
            )
            if performance
            else None
        )

        communication = (
            getattr(
                performance,
                "communication_score",
                None
            )
            if performance
            else None
        )

        compliance = (
            getattr(
                performance,
                "compliance_score",
                None
            )
            if performance
            else None
        )

        reliability = getattr(
            vendor,
            "reliability_score",
            None
        )

        overall = calculate_weighted_performance(
            performance,
            reliability
        )

        report.append(
            {
                "Vendor ID": vendor.id,

                "Vendor Name": safe_value(
                    vendor.company_name
                ),

                "Contact Person": safe_value(
                    vendor.contact_person
                ),

                "Category": safe_value(
                    vendor.category
                ),

                "Email": safe_value(
                    vendor.email
                ),

                "Delivery": format_score(
                    delivery
                ),

                "Quality": format_score(
                    quality
                ),

                "Communication": format_score(
                    communication
                ),

                "Compliance": format_score(
                    compliance
                ),

                "Reliability": format_score(
                    reliability
                ),

                "Overall Performance": format_score(
                    overall
                ),
            }
        )

    return report


# ============================================================
# VENDOR-ONLY PERFORMANCE REPORT
# ============================================================

def get_vendor_performance_for_vendor(
    db,
    vendor
):
    performance = get_latest_vendor_performance(
        db,
        vendor.id
    )

    delivery = (
        getattr(
            performance,
            "delivery_score",
            None
        )
        if performance
        else None
    )

    quality = (
        getattr(
            performance,
            "quality_score",
            None
        )
        if performance
        else None
    )

    communication = (
        getattr(
            performance,
            "communication_score",
            None
        )
        if performance
        else None
    )

    compliance = (
        getattr(
            performance,
            "compliance_score",
            None
        )
        if performance
        else None
    )

    reliability = getattr(
        vendor,
        "reliability_score",
        None
    )

    overall = calculate_weighted_performance(
        performance,
        reliability
    )

    return [
        {
            "Vendor ID": vendor.id,

            "Vendor Name": safe_value(
                vendor.company_name
            ),

            "Contact Person": safe_value(
                vendor.contact_person
            ),

            "Category": safe_value(
                vendor.category
            ),

            "Email": safe_value(
                vendor.email
            ),

            "Delivery": format_score(
                delivery
            ),

            "Quality": format_score(
                quality
            ),

            "Communication": format_score(
                communication
            ),

            "Compliance": format_score(
                compliance
            ),

            "Reliability": format_score(
                reliability
            ),

            "Overall Performance": format_score(
                overall
            ),
        }
    ]


# ============================================================
# INTERNAL PROCUREMENT REPORT
# ============================================================

def get_procurement_report(db):
    rows = (
        db.query(
            ProcurementRequest,
            Vendor
        )
        .outerjoin(
            Vendor,
            ProcurementRequest.vendor_id == Vendor.id
        )
        .order_by(
            ProcurementRequest.id.desc()
        )
        .all()
    )

    report = []

    for request, vendor in rows:
        report.append(
            {
                "Request Number": safe_value(
                    request.request_number
                ),

                "Vendor": (
                    safe_value(
                        vendor.company_name
                    )
                    if vendor
                    else "—"
                ),

                "Description": safe_value(
                    request.description
                ),

                "Quantity": safe_value(
                    request.quantity
                ),

                "Estimated Amount": format_amount(
                    request.estimated_amount
                ),

                "Expected Delivery": format_date(
                    request.expected_delivery_date
                ),

                "Status": format_status(
                    request.status
                ),

                "Created At": format_datetime(
                    request.created_at
                ),
            }
        )

    return report


# ============================================================
# VENDOR-ONLY PROCUREMENT REPORT
# ============================================================

def get_vendor_procurement_report(
    db,
    vendor
):
    rows = (
        db.query(ProcurementRequest)
        .filter(
            ProcurementRequest.vendor_id == vendor.id
        )
        .order_by(
            ProcurementRequest.id.desc()
        )
        .all()
    )

    report = []

    for request in rows:
        report.append(
            {
                "Request Number": safe_value(
                    request.request_number
                ),

                "Vendor": safe_value(
                    vendor.company_name
                ),

                "Description": safe_value(
                    request.description
                ),

                "Quantity": safe_value(
                    request.quantity
                ),

                "Estimated Amount": format_amount(
                    request.estimated_amount
                ),

                "Expected Delivery": format_date(
                    request.expected_delivery_date
                ),

                "Status": format_status(
                    request.status
                ),

                "Created At": format_datetime(
                    request.created_at
                ),
            }
        )

    return report


# ============================================================
# INTERNAL PURCHASE ORDER REPORT
# ============================================================

def get_purchase_order_report(db):
    rows = (
        db.query(
            PurchaseOrder,
            Vendor,
            Supplier
        )
        .outerjoin(
            Vendor,
            PurchaseOrder.vendor_id == Vendor.id
        )
        .outerjoin(
            Supplier,
            PurchaseOrder.supplier_id == Supplier.id
        )
        .order_by(
            PurchaseOrder.id.desc()
        )
        .all()
    )

    report = []

    for order, vendor, supplier in rows:
        report.append(
            {
                "PO Number": safe_value(
                    order.order_number
                ),

                "Vendor": (
                    safe_value(
                        vendor.company_name
                    )
                    if vendor
                    else "—"
                ),

                "Supplier": (
                    safe_value(
                        supplier.company_name
                    )
                    if supplier
                    else "—"
                ),

                "Order Date": format_date(
                    order.order_date
                ),

                "Expected Delivery": format_date(
                    order.expected_delivery_date
                ),

                "Actual Delivery": format_date(
                    order.actual_delivery_date
                ),

                "Total Amount": format_amount(
                    order.total_amount
                ),

                "Status": format_status(
                    order.status
                ),
            }
        )

    return report


# ============================================================
# VENDOR-ONLY PURCHASE ORDER REPORT
# ============================================================

def get_vendor_purchase_order_report(
    db,
    vendor
):
    rows = (
        db.query(PurchaseOrder)
        .filter(
            PurchaseOrder.vendor_id == vendor.id
        )
        .order_by(
            PurchaseOrder.id.desc()
        )
        .all()
    )

    report = []

    for order in rows:
        supplier = None

        if order.supplier_id:
            supplier = (
                db.query(Supplier)
                .filter(
                    Supplier.id == order.supplier_id
                )
                .first()
            )

        report.append(
            {
                "PO Number": safe_value(
                    order.order_number
                ),

                "Vendor": safe_value(
                    vendor.company_name
                ),

                "Supplier": (
                    safe_value(
                        supplier.company_name
                    )
                    if supplier
                    else "—"
                ),

                "Order Date": format_date(
                    order.order_date
                ),

                "Expected Delivery": format_date(
                    order.expected_delivery_date
                ),

                "Actual Delivery": format_date(
                    order.actual_delivery_date
                ),

                "Total Amount": format_amount(
                    order.total_amount
                ),

                "Status": format_status(
                    order.status
                ),
            }
        )

    return report


# ============================================================
# INTERNAL COMPLIANCE REPORT
# ============================================================

def get_compliance_report(db):
    vendors = (
        db.query(Vendor)
        .order_by(
            Vendor.id.asc()
        )
        .all()
    )

    report = []

    for vendor in vendors:
        report.append(
            {
                "Vendor ID": vendor.id,

                "Vendor Name": safe_value(
                    vendor.company_name
                ),

                "Category": safe_value(
                    vendor.category
                ),

                "Email": safe_value(
                    vendor.email
                ),

                "Vendor Status": format_status(
                    getattr(
                        vendor,
                        "vendor_status",
                        None
                    )
                ),

                "Approval Status": format_status(
                    getattr(
                        vendor,
                        "approval_status",
                        None
                    )
                ),
            }
        )

    return report


# ============================================================
# CONTRACT STATUS
# ============================================================

def calculate_contract_status(
    end_date,
    current_status=None
):
    status = (
        str(current_status).lower()
        if current_status
        else None
    )

    if status in [
        "terminated",
        "draft"
    ]:
        return format_status(status)

    if not end_date:
        return (
            format_status(current_status)
            if current_status
            else "Draft"
        )

    today = datetime.now().date()

    try:
        end = end_date.date()
    except Exception:
        end = end_date

    if end < today:
        return "Expired"

    days_left = (
        end - today
    ).days

    if days_left <= 30:
        return "Expiring"

    return "Active"


# ============================================================
# INTERNAL CONTRACT REPORT
# ============================================================

def get_contract_report(db):
    rows = (
        db.query(
            Contract,
            Vendor
        )
        .outerjoin(
            Vendor,
            Contract.vendor_id == Vendor.id
        )
        .order_by(
            Contract.id.desc()
        )
        .all()
    )

    report = []

    for contract, vendor in rows:
        status = calculate_contract_status(
            getattr(
                contract,
                "end_date",
                None
            ),
            getattr(
                contract,
                "status",
                None
            )
        )

        report.append(
            {
                "Contract Number": safe_value(
                    getattr(
                        contract,
                        "contract_number",
                        None
                    )
                ),

                "Vendor": (
                    safe_value(
                        vendor.company_name
                    )
                    if vendor
                    else "—"
                ),

                "Title": safe_value(
                    getattr(
                        contract,
                        "title",
                        None
                    )
                ),

                "Start Date": format_date(
                    getattr(
                        contract,
                        "start_date",
                        None
                    )
                ),

                "End Date": format_date(
                    getattr(
                        contract,
                        "end_date",
                        None
                    )
                ),

                "Contract Value": format_amount(
                    getattr(
                        contract,
                        "contract_value",
                        None
                    )
                ),

                "Status": status,
            }
        )

    return report


# ============================================================
# VENDOR-ONLY CONTRACT REPORT
# ============================================================

def get_vendor_contract_report(
    db,
    vendor
):
    rows = (
        db.query(Contract)
        .filter(
            Contract.vendor_id == vendor.id
        )
        .order_by(
            Contract.id.desc()
        )
        .all()
    )

    report = []

    for contract in rows:
        status = calculate_contract_status(
            getattr(
                contract,
                "end_date",
                None
            ),
            getattr(
                contract,
                "status",
                None
            )
        )

        report.append(
            {
                "Contract Number": safe_value(
                    getattr(
                        contract,
                        "contract_number",
                        None
                    )
                ),

                "Vendor": safe_value(
                    vendor.company_name
                ),

                "Title": safe_value(
                    getattr(
                        contract,
                        "title",
                        None
                    )
                ),

                "Start Date": format_date(
                    getattr(
                        contract,
                        "start_date",
                        None
                    )
                ),

                "End Date": format_date(
                    getattr(
                        contract,
                        "end_date",
                        None
                    )
                ),

                "Contract Value": format_amount(
                    getattr(
                        contract,
                        "contract_value",
                        None
                    )
                ),

                "Status": status,
            }
        )

    return report


# ============================================================
# INTERNAL INVOICE REPORT
# ============================================================

def get_invoice_report(db):
    rows = (
        db.query(
            Invoice,
            Vendor,
            PurchaseOrder
        )
        .outerjoin(
            Vendor,
            Invoice.vendor_id == Vendor.id
        )
        .outerjoin(
            PurchaseOrder,
            Invoice.purchase_order_id == PurchaseOrder.id
        )
        .order_by(
            Invoice.id.desc()
        )
        .all()
    )

    report = []

    for invoice, vendor, purchase_order in rows:
        report.append(
            {
                "Invoice Number": safe_value(
                    invoice.invoice_number
                ),

                "Vendor": (
                    safe_value(
                        vendor.company_name
                    )
                    if vendor
                    else "—"
                ),

                "PO Number": (
                    safe_value(
                        purchase_order.order_number
                    )
                    if purchase_order
                    else "—"
                ),

                "Invoice Date": format_date(
                    invoice.invoice_date
                ),

                "Due Date": format_date(
                    invoice.due_date
                ),

                "Amount": format_amount(
                    invoice.amount
                ),

                "Status": format_status(
                    invoice.status
                ),

                "Verified At": format_datetime(
                    invoice.verified_at
                ),

                "Paid At": format_datetime(
                    invoice.paid_at
                ),
            }
        )

    return report


# ============================================================
# VENDOR-ONLY INVOICE REPORT
# ============================================================

def get_vendor_invoice_report(
    db,
    vendor
):
    rows = (
        db.query(Invoice)
        .filter(
            Invoice.vendor_id == vendor.id
        )
        .order_by(
            Invoice.id.desc()
        )
        .all()
    )

    report = []

    for invoice in rows:
        purchase_order = None

        if invoice.purchase_order_id:
            purchase_order = (
                db.query(PurchaseOrder)
                .filter(
                    PurchaseOrder.id == invoice.purchase_order_id
                )
                .first()
            )

        report.append(
            {
                "Invoice Number": safe_value(
                    invoice.invoice_number
                ),

                "Vendor": safe_value(
                    vendor.company_name
                ),

                "PO Number": (
                    safe_value(
                        purchase_order.order_number
                    )
                    if purchase_order
                    else "—"
                ),

                "Invoice Date": format_date(
                    invoice.invoice_date
                ),

                "Due Date": format_date(
                    invoice.due_date
                ),

                "Amount": format_amount(
                    invoice.amount
                ),

                "Status": format_status(
                    invoice.status
                ),

                "Verified At": format_datetime(
                    invoice.verified_at
                ),

                "Paid At": format_datetime(
                    invoice.paid_at
                ),
            }
        )

    return report


# ============================================================
# INTERNAL PAYMENT REPORT
# ============================================================

def get_payment_report(db):
    rows = (
        db.query(
            Payment,
            Invoice,
            Vendor
        )
        .outerjoin(
            Invoice,
            Payment.invoice_id == Invoice.id
        )
        .outerjoin(
            Vendor,
            Payment.vendor_id == Vendor.id
        )
        .order_by(
            Payment.id.desc()
        )
        .all()
    )

    report = []

    for payment, invoice, vendor in rows:
        report.append(
            {
                "Payment Number": safe_value(
                    payment.payment_number
                ),

                "Invoice Number": (
                    safe_value(
                        invoice.invoice_number
                    )
                    if invoice
                    else "—"
                ),

                "Vendor": (
                    safe_value(
                        vendor.company_name
                    )
                    if vendor
                    else "—"
                ),

                "Amount": format_amount(
                    payment.amount
                ),

                "Payment Date": format_date(
                    payment.payment_date
                ),

                "Payment Method": safe_value(
                    payment.payment_method
                ),

                "Transaction Reference": safe_value(
                    payment.transaction_reference
                ),

                "Status": format_status(
                    payment.status
                ),
            }
        )

    return report


# ============================================================
# VENDOR-ONLY PAYMENT REPORT
# ============================================================

def get_vendor_payment_report(
    db,
    vendor
):
    rows = (
        db.query(Payment)
        .filter(
            Payment.vendor_id == vendor.id
        )
        .order_by(
            Payment.id.desc()
        )
        .all()
    )

    report = []

    for payment in rows:
        invoice = None

        if payment.invoice_id:
            invoice = (
                db.query(Invoice)
                .filter(
                    Invoice.id == payment.invoice_id
                )
                .first()
            )

        report.append(
            {
                "Payment Number": safe_value(
                    payment.payment_number
                ),

                "Invoice Number": (
                    safe_value(
                        invoice.invoice_number
                    )
                    if invoice
                    else "—"
                ),

                "Vendor": safe_value(
                    vendor.company_name
                ),

                "Amount": format_amount(
                    payment.amount
                ),

                "Payment Date": format_date(
                    payment.payment_date
                ),

                "Payment Method": safe_value(
                    payment.payment_method
                ),

                "Transaction Reference": safe_value(
                    payment.transaction_reference
                ),

                "Status": format_status(
                    payment.status
                ),
            }
        )

    return report


# ============================================================
# REPORT DATA DISPATCHER
# ============================================================

def get_report_data(
    report_type,
    db,
    current_user
):
    check_report_access(
        current_user,
        report_type
    )

    role = str(current_user.role)

    # ========================================================
    # VENDOR
    # ========================================================

    if role == "vendor":
        vendor = get_current_vendor(
            db,
            current_user
        )

        if report_type == "procurement":
            return get_vendor_procurement_report(
                db,
                vendor
            )

        if report_type == "purchase-orders":
            return get_vendor_purchase_order_report(
                db,
                vendor
            )

        if report_type == "vendor-performance":
            return get_vendor_performance_for_vendor(
                db,
                vendor
            )

        if report_type == "contracts":
            return get_vendor_contract_report(
                db,
                vendor
            )

        if report_type == "invoices":
            return get_vendor_invoice_report(
                db,
                vendor
            )

        if report_type == "payments":
            return get_vendor_payment_report(
                db,
                vendor
            )

    # ========================================================
    # INTERNAL USERS
    # ========================================================

    if report_type == "procurement":
        return get_procurement_report(db)

    if report_type == "purchase-orders":
        return get_purchase_order_report(db)

    if report_type == "vendor-performance":
        return get_vendor_performance_report(db)

    if report_type == "compliance":
        return get_compliance_report(db)

    if report_type == "contracts":
        return get_contract_report(db)

    if report_type == "invoices":
        return get_invoice_report(db)

    if report_type == "payments":
        return get_payment_report(db)

    raise HTTPException(
        status_code=404,
        detail="Report type not found"
    )


# ============================================================
# REPORT TITLES
# ============================================================

REPORT_TITLES = {
    "procurement": "Procurement Report",
    "purchase-orders": "Purchase Order Report",
    "vendor-performance": "Vendor Performance Report",
    "compliance": "Compliance Report",
    "contracts": "Contract Report",
    "invoices": "Invoice Report",
    "payments": "Payment Report",
}


# ============================================================
# DATA ENDPOINT
# ============================================================

@router.get("/{report_type}/data")
def get_report_data_endpoint(
    report_type: str,
    db: Session = Depends(get_db),
    current_user=Depends(get_current_user)
):
    data = get_report_data(
        report_type,
        db,
        current_user
    )

    return {
        "report_type": report_type,
        "title": REPORT_TITLES.get(
            report_type,
            "Report"
        ),
        "record_count": len(data),
        "data": data,
    }


# ============================================================
# AVAILABLE REPORTS
# ============================================================

@router.get("/available")
def get_available_reports(
    current_user=Depends(get_current_user)
):
    check_report_access(
        current_user
    )

    return {
        "role": str(current_user.role),
        "reports": get_allowed_reports(
            current_user
        )
    }


# ============================================================
# PDF GENERATION
# ============================================================

def generate_pdf(
    report_type,
    data
):
    buffer = BytesIO()

    document = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        rightMargin=20,
        leftMargin=20,
        topMargin=25,
        bottomMargin=25,
    )

    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        "ReportTitle",
        parent=styles["Title"],
        alignment=TA_CENTER,
        fontSize=18,
        leading=22,
        spaceAfter=8,
    )

    info_style = ParagraphStyle(
        "ReportInfo",
        parent=styles["Normal"],
        alignment=TA_CENTER,
        fontSize=9,
        leading=12,
        spaceAfter=15,
    )

    cell_style = ParagraphStyle(
        "TableCell",
        parent=styles["Normal"],
        fontSize=7,
        leading=9,
    )

    header_style = ParagraphStyle(
        "TableHeader",
        parent=styles["Normal"],
        fontSize=7,
        leading=9,
        alignment=TA_CENTER,
    )

    story = []

    title = REPORT_TITLES.get(
        report_type,
        "Report"
    )

    story.append(
        Paragraph(
            title,
            title_style
        )
    )

    story.append(
        Paragraph(
            f"Generated on "
            f"{datetime.now().strftime('%d %b %Y %H:%M')}"
            f" &nbsp;&nbsp;|&nbsp;&nbsp; "
            f"Records: {len(data)}",
            info_style
        )
    )

    if not data:
        story.append(
            Paragraph(
                "No records available.",
                styles["Normal"]
            )
        )

        document.build(story)

        buffer.seek(0)

        return buffer

    headers = list(
        data[0].keys()
    )

    table_data = [
        [
            Paragraph(
                safe_value(header),
                header_style
            )
            for header in headers
        ]
    ]

    for row in data:
        table_data.append(
            [
                Paragraph(
                    safe_value(
                        row.get(header)
                    ),
                    cell_style
                )
                for header in headers
            ]
        )

    # ========================================================
    # PDF COLUMN WIDTHS
    # ========================================================

    width_maps = {
        "procurement": {
            "Request Number": 0.90,
            "Vendor": 1.35,
            "Description": 2.00,
            "Quantity": 0.65,
            "Estimated Amount": 1.10,
            "Expected Delivery": 1.10,
            "Status": 0.95,
            "Created At": 1.25,
        },

        "purchase-orders": {
            "PO Number": 0.95,
            "Vendor": 1.40,
            "Supplier": 1.35,
            "Order Date": 1.00,
            "Expected Delivery": 1.10,
            "Actual Delivery": 1.10,
            "Total Amount": 1.20,
            "Status": 0.95,
        },

        "vendor-performance": {
            "Vendor ID": 0.55,
            "Vendor Name": 1.35,
            "Contact Person": 1.15,
            "Category": 1.00,
            "Email": 1.55,
            "Delivery": 0.75,
            "Quality": 0.75,
            "Communication": 0.85,
            "Compliance": 0.75,
            "Reliability": 0.80,
            "Overall Performance": 1.05,
        },

        "compliance": {
            "Vendor ID": 0.70,
            "Vendor Name": 1.55,
            "Category": 1.20,
            "Email": 1.80,
            "Vendor Status": 1.10,
            "Approval Status": 1.20,
        },

        "contracts": {
            "Contract Number": 1.15,
            "Vendor": 1.55,
            "Title": 2.00,
            "Start Date": 1.10,
            "End Date": 1.10,
            "Contract Value": 1.25,
            "Status": 1.00,
        },

        "invoices": {
            "Invoice Number": 1.15,
            "Vendor": 1.45,
            "PO Number": 1.05,
            "Invoice Date": 1.05,
            "Due Date": 1.05,
            "Amount": 1.10,
            "Status": 1.00,
            "Verified At": 1.25,
            "Paid At": 1.25,
        },

        "payments": {
            "Payment Number": 1.10,
            "Invoice Number": 1.15,
            "Vendor": 1.40,
            "Amount": 1.10,
            "Payment Date": 1.05,
            "Payment Method": 1.15,
            "Transaction Reference": 1.70,
            "Status": 1.00,
        },
    }

    widths_map = width_maps.get(
        report_type,
        {}
    )

    col_widths = []

    for header in headers:
        width = widths_map.get(
            header,
            1.0
        )

        col_widths.append(
            width * inch
        )

    table = Table(
        table_data,
        colWidths=col_widths,
        repeatRows=1,
        hAlign="CENTER",
    )

    table.setStyle(
        TableStyle(
            [
                (
                    "BACKGROUND",
                    (0, 0),
                    (-1, 0),
                    colors.HexColor("#E5E7EB"),
                ),

                (
                    "TEXTCOLOR",
                    (0, 0),
                    (-1, 0),
                    colors.black,
                ),

                (
                    "FONTNAME",
                    (0, 0),
                    (-1, 0),
                    "Helvetica-Bold",
                ),

                (
                    "GRID",
                    (0, 0),
                    (-1, -1),
                    0.4,
                    colors.grey,
                ),

                (
                    "VALIGN",
                    (0, 0),
                    (-1, -1),
                    "MIDDLE",
                ),

                (
                    "ALIGN",
                    (0, 0),
                    (-1, 0),
                    "CENTER",
                ),

                (
                    "LEFTPADDING",
                    (0, 0),
                    (-1, -1),
                    4,
                ),

                (
                    "RIGHTPADDING",
                    (0, 0),
                    (-1, -1),
                    4,
                ),

                (
                    "TOPPADDING",
                    (0, 0),
                    (-1, -1),
                    4,
                ),

                (
                    "BOTTOMPADDING",
                    (0, 0),
                    (-1, -1),
                    4,
                ),
            ]
        )
    )

    story.append(table)

    document.build(story)

    buffer.seek(0)

    return buffer


# ============================================================
# PDF ENDPOINT
# ============================================================

@router.get("/{report_type}/pdf")
def download_report_pdf(
    report_type: str,
    db: Session = Depends(get_db),
    current_user=Depends(get_current_user)
):
    data = get_report_data(
        report_type,
        db,
        current_user
    )

    buffer = generate_pdf(
        report_type,
        data
    )

    filename = f"{report_type}-report.pdf"

    return StreamingResponse(
        buffer,
        media_type="application/pdf",
        headers={
            "Content-Disposition":
                f'inline; filename="{filename}"'
        }
    )


# ============================================================
# EXCEL GENERATION
# ============================================================

def generate_excel(
    report_type,
    data
):
    workbook = Workbook()

    worksheet = workbook.active

    worksheet.title = (
        REPORT_TITLES.get(
            report_type,
            "Report"
        )[:31]
    )

    if not data:
        worksheet["A1"] = (
            REPORT_TITLES.get(
                report_type,
                "Report"
            )
        )

        worksheet["A2"] = (
            "No records available."
        )

        buffer = BytesIO()

        workbook.save(buffer)

        buffer.seek(0)

        return buffer

    headers = list(
        data[0].keys()
    )

    # ========================================================
    # HEADER ROW
    # ========================================================

    for column_index, header in enumerate(
        headers,
        start=1
    ):
        cell = worksheet.cell(
            row=1,
            column=column_index,
            value=header
        )

        cell.font = Font(
            bold=True
        )

        cell.fill = PatternFill(
            fill_type="solid",
            fgColor="E5E7EB"
        )

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )

    # ========================================================
    # DATA ROWS
    # ========================================================

    for row_index, row in enumerate(
        data,
        start=2
    ):
        for column_index, header in enumerate(
            headers,
            start=1
        ):
            value = row.get(
                header
            )

            worksheet.cell(
                row=row_index,
                column=column_index,
                value=safe_value(value)
            )

    # ========================================================
    # FREEZE HEADER
    # ========================================================

    worksheet.freeze_panes = "A2"

    # ========================================================
    # AUTO WIDTH
    # ========================================================

    for column_index, header in enumerate(
        headers,
        start=1
    ):
        max_length = len(
            str(header)
        )

        for row_index in range(
            2,
            worksheet.max_row + 1
        ):
            value = worksheet.cell(
                row=row_index,
                column=column_index
            ).value

            if value is not None:
                max_length = max(
                    max_length,
                    len(str(value))
                )

        width = min(
            max_length + 2,
            45
        )

        worksheet.column_dimensions[
            get_column_letter(
                column_index
            )
        ].width = width

    buffer = BytesIO()

    workbook.save(buffer)

    buffer.seek(0)

    return buffer


# ============================================================
# EXCEL ENDPOINT
# ============================================================

@router.get("/{report_type}/excel")
def download_report_excel(
    report_type: str,
    db: Session = Depends(get_db),
    current_user=Depends(get_current_user)
):
    data = get_report_data(
        report_type,
        db,
        current_user
    )

    buffer = generate_excel(
        report_type,
        data
    )

    filename = f"{report_type}-report.xlsx"

    return StreamingResponse(
        buffer,
        media_type=(
            "application/vnd.openxmlformats-"
            "officedocument.spreadsheetml.sheet"
        ),
        headers={
            "Content-Disposition":
                f'attachment; filename="{filename}"'
        }
    )