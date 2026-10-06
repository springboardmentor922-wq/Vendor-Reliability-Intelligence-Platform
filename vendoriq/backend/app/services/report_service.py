"""Reports & Export module."""

import os
import uuid
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font
from reportlab.lib import colors
from reportlab.lib.pagesizes import landscape, A4
from reportlab.platypus import (
    SimpleDocTemplate,
    Table,
    TableStyle,
    Paragraph,
    Spacer,
)
from reportlab.lib.styles import getSampleStyleSheet
from sqlalchemy.orm import Session

from app.core.config import settings
from app.models.vendor import Vendor
from app.models.purchase_order import PurchaseOrder, POStatus
from app.models.procurement import ProcurementRequest
from app.models.contract import Contract, Certification
from app.services import performance_service


REPORTS_DIR = os.path.abspath(
    os.path.join(settings.UPLOAD_DIR, "reports")
)
os.makedirs(REPORTS_DIR, exist_ok=True)


def _filename(report_type: str, ext: str) -> str:
    ts = datetime.utcnow().strftime("%Y%m%d%H%M%S")
    return os.path.join(
        REPORTS_DIR,
        f"{report_type}_{ts}_{uuid.uuid4().hex[:8]}.{ext}",
    )


def build_excel(
    report_type: str,
    title: str,
    headers: list[str],
    rows: list[list],
) -> str:
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]

    ws.append([title])
    ws["A1"].font = Font(bold=True, size=14)

    ws.append([
        f"Generated {datetime.utcnow().strftime('%Y-%m-%d %H:%M UTC')}"
    ])
    ws.append([])

    ws.append(headers)

    for cell in ws[4]:
        cell.font = Font(bold=True)

    for row in rows:
        ws.append(row)

    for col in ws.columns:
        max_len = max(
            (len(str(c.value)) for c in col if c.value is not None),
            default=10,
        )
        ws.column_dimensions[
            col[0].column_letter
        ].width = min(max_len + 4, 45)

    path = _filename(report_type, "xlsx")
    wb.save(path)

    return path


def build_pdf(
    report_type: str,
    title: str,
    headers: list[str],
    rows: list[list],
) -> str:
    path = _filename(report_type, "pdf")

    doc = SimpleDocTemplate(
        path,
        pagesize=landscape(A4),
    )

    styles = getSampleStyleSheet()

    elements = [
        Paragraph(title, styles["Title"]),
        Paragraph(
            f"Generated {datetime.utcnow().strftime('%Y-%m-%d %H:%M UTC')}",
            styles["Normal"],
        ),
        Spacer(1, 12),
    ]

    data = [headers] + [
        [str(value) for value in row]
        for row in rows
    ]

    table = Table(data, repeatRows=1)

    table.setStyle(
        TableStyle(
            [
                (
                    "BACKGROUND",
                    (0, 0),
                    (-1, 0),
                    colors.HexColor("#1e3a8a"),
                ),
                (
                    "TEXTCOLOR",
                    (0, 0),
                    (-1, 0),
                    colors.white,
                ),
                (
                    "FONTNAME",
                    (0, 0),
                    (-1, 0),
                    "Helvetica-Bold",
                ),
                (
                    "FONTSIZE",
                    (0, 0),
                    (-1, -1),
                    8,
                ),
                (
                    "GRID",
                    (0, 0),
                    (-1, -1),
                    0.5,
                    colors.grey,
                ),
                (
                    "ROWBACKGROUNDS",
                    (0, 1),
                    (-1, -1),
                    [
                        colors.white,
                        colors.HexColor("#f4f6f9"),
                    ],
                ),
                (
                    "VALIGN",
                    (0, 0),
                    (-1, -1),
                    "TOP",
                ),
            ]
        )
    )

    elements.append(table)
    doc.build(elements)

    return path


def _export(
    report_type: str,
    title: str,
    headers: list[str],
    rows: list[list],
    report_format: str,
):
    if report_format == "xlsx":
        return build_excel(
            report_type,
            title,
            headers,
            rows,
        )

    if report_format == "pdf":
        return build_pdf(
            report_type,
            title,
            headers,
            rows,
        )

    return {
        "title": title,
        "headers": headers,
        "rows": rows,
    }


def generate_vendor_performance_report(
    db: Session,
    vendor_id=None,
    format: str = "xlsx",
):
    vendors = (
        [db.query(Vendor).filter(Vendor.id == vendor_id).first()]
        if vendor_id
        else db.query(Vendor).all()
    )

    vendors = [v for v in vendors if v is not None]

    headers = [
        "Vendor",
        "Category",
        "Status",
        "On-Time Deliveries",
        "Delayed Deliveries",
        "Delivery Rate",
        "Completed Orders",
        "Completion Rate",
        "Rating",
    ]

    rows = []

    for vendor in vendors:
        on_time, delayed = (
            performance_service.compute_delivery_metrics(
                db,
                vendor.id,
            )
        )

        total_deliveries = on_time + delayed

        delivery_rate = (
            round(
                (on_time / total_deliveries) * 100,
                2,
            )
            if total_deliveries
            else 0.0
        )

        total_orders = (
            db.query(PurchaseOrder)
            .filter(
                PurchaseOrder.vendor_id == vendor.id,
                PurchaseOrder.status != POStatus.CANCELLED,
            )
            .count()
        )

        completed_orders = (
            db.query(PurchaseOrder)
            .filter(
                PurchaseOrder.vendor_id == vendor.id,
                PurchaseOrder.status == POStatus.COMPLETED,
            )
            .count()
        )

        completion_rate = (
            round(
                (completed_orders / total_orders) * 100,
                2,
            )
            if total_orders
            else 0.0
        )

        rows.append(
            [
                vendor.company_name,
                (
                    vendor.category.value
                    if vendor.category
                    else "-"
                ),
                (
                    vendor.status.value
                    if vendor.status
                    else "-"
                ),
                on_time,
                delayed,
                delivery_rate,
                completed_orders,
                completion_rate,
                round(vendor.rating or 0, 2),
            ]
        )

    return _export(
        "vendor_performance",
        "Vendor Performance Report",
        headers,
        rows,
        format,
    )


def generate_procurement_report(
    db: Session,
    format: str = "xlsx",
):
    requests = (
        db.query(ProcurementRequest)
        .order_by(
            ProcurementRequest.created_at.desc()
        )
        .all()
    )

    headers = [
        "Request Number",
        "Title",
        "Department",
        "Category",
        "Quantity",
        "Budget",
        "Priority",
        "Status",
        "Required Date",
        "Created At",
    ]

    rows = []

    for request in requests:
        rows.append(
            [
                request.request_number,
                request.title,
                request.department or "-",
                request.category or "-",
                request.quantity,
                round(
                    request.estimated_budget or 0,
                    2,
                ),
                (
                    request.priority.value
                    if request.priority
                    else "-"
                ),
                (
                    request.status.value
                    if request.status
                    else "-"
                ),
                (
                    request.required_date.strftime(
                        "%Y-%m-%d"
                    )
                    if request.required_date
                    else "-"
                ),
                (
                    request.created_at.strftime(
                        "%Y-%m-%d %H:%M"
                    )
                    if request.created_at
                    else "-"
                ),
            ]
        )

    return _export(
        "procurement",
        "Procurement Report",
        headers,
        rows,
        format,
    )


def generate_purchase_order_report(
    db: Session,
    vendor_id=None,
    format: str = "xlsx",
):
    query = db.query(PurchaseOrder)

    if vendor_id:
        query = query.filter(
            PurchaseOrder.vendor_id == vendor_id
        )

    orders = (
        query.order_by(
            PurchaseOrder.created_at.desc()
        )
        .all()
    )

    headers = [
        "PO Number",
        "Vendor",
        "Total Amount",
        "Status",
        "Order Date",
        "Expected Delivery",
        "Actual Delivery",
    ]

    rows = []

    for po in orders:
        vendor = (
            db.query(Vendor)
            .filter(Vendor.id == po.vendor_id)
            .first()
        )

        rows.append(
            [
                po.po_number,
                (
                    vendor.company_name
                    if vendor
                    else str(po.vendor_id)
                ),
                round(po.total_amount or 0, 2),
                (
                    po.status.value
                    if po.status
                    else "-"
                ),
                (
                    po.order_date.strftime(
                        "%Y-%m-%d"
                    )
                    if po.order_date
                    else "-"
                ),
                (
                    po.expected_delivery_date.strftime(
                        "%Y-%m-%d"
                    )
                    if po.expected_delivery_date
                    else "-"
                ),
                (
                    po.actual_delivery_date.strftime(
                        "%Y-%m-%d"
                    )
                    if po.actual_delivery_date
                    else "-"
                ),
            ]
        )

    return _export(
        "purchase_orders",
        "Purchase Order Report",
        headers,
        rows,
        format,
    )


def generate_compliance_report(
    db: Session,
    vendor_id=None,
    format: str = "xlsx",
):
    query = db.query(Certification)

    if vendor_id:
        query = query.filter(
            Certification.vendor_id == vendor_id
        )

    certifications = query.order_by(
        Certification.expiry_date.asc()
    ).all()

    headers = [
        "Vendor",
        "Certification",
        "Issuing Body",
        "Issue Date",
        "Expiry Date",
        "Status",
    ]

    rows = []

    for cert in certifications:
        vendor = (
            db.query(Vendor)
            .filter(Vendor.id == cert.vendor_id)
            .first()
        )

        rows.append(
            [
                (
                    vendor.company_name
                    if vendor
                    else str(cert.vendor_id)
                ),
                cert.name,
                cert.issuing_body or "-",
                (
                    cert.issue_date.strftime(
                        "%Y-%m-%d"
                    )
                    if cert.issue_date
                    else "-"
                ),
                (
                    cert.expiry_date.strftime(
                        "%Y-%m-%d"
                    )
                    if cert.expiry_date
                    else "-"
                ),
                (
                    cert.status.value
                    if cert.status
                    else "-"
                ),
            ]
        )

    return _export(
        "compliance",
        "Compliance Report",
        headers,
        rows,
        format,
    )


def generate_contract_report(
    db: Session,
    vendor_id=None,
    format: str = "xlsx",
):
    query = db.query(Contract)

    if vendor_id:
        query = query.filter(
            Contract.vendor_id == vendor_id
        )

    contracts = query.order_by(
        Contract.end_date.asc()
    ).all()

    headers = [
        "Contract Number",
        "Vendor",
        "Title",
        "Start Date",
        "End Date",
        "Value",
        "Status",
    ]

    rows = []

    for contract in contracts:
        vendor = (
            db.query(Vendor)
            .filter(Vendor.id == contract.vendor_id)
            .first()
        )

        rows.append(
            [
                contract.contract_number,
                (
                    vendor.company_name
                    if vendor
                    else str(contract.vendor_id)
                ),
                contract.title,
                (
                    contract.start_date.strftime(
                        "%Y-%m-%d"
                    )
                    if contract.start_date
                    else "-"
                ),
                (
                    contract.end_date.strftime(
                        "%Y-%m-%d"
                    )
                    if contract.end_date
                    else "-"
                ),
                round(contract.value or 0, 2),
                (
                    contract.status.value
                    if contract.status
                    else "-"
                ),
            ]
        )

    return _export(
        "contracts",
        "Contract Report",
        headers,
        rows,
        format,
    )